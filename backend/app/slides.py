"""Five deterministic 16:9 layouts. Only validated content enters the renderer."""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
import math
import unicodedata
from .core import AppError

NAVY, PURPLE, INK, MUTED, WHITE = '101827', '6658E8', '182235', '6B7280', 'FFFFFF'


def text(slide, value, x, y, width, height, size=22, color=INK, bold=False):
    lines = value.split('\n')
    # Estimate CJK and Latin widths conservatively before placing any text.
    def estimated_height(font_size):
        usable = max(1, width * 72 - 12)
        count = sum(max(1, math.ceil(sum(1 if unicodedata.east_asian_width(c) in 'WF' else .58 for c in line) * font_size / usable)) for line in lines)
        return count * font_size * 1.3 + max(0, len(lines) - 1) * 12
    while size > 10 and estimated_height(size) > height * 72 - 8:
        size -= 1
    if estimated_height(size) > height * 72 - 8:
        raise AppError('PPT 文本超出版式容量，请减少内容后重试', 502)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.06)
    frame.margin_top = frame.margin_bottom = Inches(.04)
    for index, line in enumerate(lines):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = 'Microsoft YaHei'
        p.font.size, p.font.bold = Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(12 if index < len(lines) - 1 else 0)
    return box


def render_deck(deck, path):
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    for index, spec in enumerate(deck.slides):
        if any(len(item) > 64 for item in [*spec.left, *spec.right]):
            raise AppError('PPT 要点过长，已停止生成以避免文字溢出', 502)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        dark = spec.layout in ('cover', 'section')
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(NAVY if dark else 'F7F8FC')
        title_color = WHITE if dark else INK
        text(slide, 'CORTEX  /  SECOND BRAIN', .7, .35, 10, .4, 11, PURPLE, True)
        text(slide, f'{index+1:02d} / {len(deck.slides):02d}', 11.65, 6.95, 1.1, .3, 10, MUTED)
        if dark:
            text(slide, spec.title, .85, 2.0, 11.5, 2, 38, title_color, True)
            text(slide, spec.subtitle, .85, 4.3, 10.8, 1.5, 20, 'BDC5D5')
        else:
            text(slide, spec.title, .7, 1.0, 11.9, 1.0, 30, title_color, True)
            text(slide, spec.subtitle, .7, 2.0, 11.9, .65, 16, MUTED)
            if spec.layout == 'chart':
                if not spec.categories or len(spec.categories) != len(spec.values) or not spec.data_note:
                    raise AppError('图表缺少有效数据或数据说明', 502)
                if any(len(label) > 18 for label in spec.categories):
                    raise AppError('图表标签过长', 502)
                data = CategoryChartData()
                data.categories = spec.categories
                data.add_series('数据', spec.values)
                chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.85), Inches(2.8), Inches(11.5), Inches(3.5), data).chart
                chart.has_legend = False
                text(slide, spec.data_note, .85, 6.4, 11.5, .45, 11, MUTED)
            else:
                text(slide, '\n'.join('• ' + item for item in spec.left), .85, 2.85, 5.6, 3.7, 19)
                text(slide, '\n'.join('• ' + item for item in spec.right), 6.9, 2.85, 5.5, 3.7, 19)
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)
